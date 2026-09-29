import sys; sys.path.insert(0, "/tmp/diag")
from layout import TITLE, SUBTITLE
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

d = docx.Document()
s = d.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = Inches(11.69), Inches(8.27)      # A4 landscape
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(s, m, Inches(0.6))

def para(text, size=11, bold=False, colour=None, after=6, align=None):
    p = d.add_paragraph()
    if align: p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold
    r.font.name = "Calibri"
    if colour: r.font.color.rgb = RGBColor.from_string(colour)
    p.paragraph_format.space_after = Pt(after)
    return p

para("MEA Orchestration — data flow", 20, True, "1F2328", after=2)
para("Ben-Shalom Lab · how a recording becomes a report", 11, False, "6B7684", after=14)

d.add_picture("/tmp/diag/orchestrator_dataflow.png", width=Inches(10.4))
d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

d.add_page_break()
para("What each stage does", 15, True, "1F2328", after=10)

NOTES = [
    ("Instruments → NAS",
     "MaxTwo and MaxOne write raw HDF5 to the recording computers, which copy each "
     "session to the Ben-Shalom NAS under <date>/<chip>/<assay>/<run>/data.raw.h5."),
    ("Watcher",
     "Polls the NAS for recordings that have finished copying. A folder counts as "
     "complete only when MaxWell's own finished marker is present and the folder has "
     "stopped changing for a set time — starting earlier would analyse a half-copied file."),
    ("Queue",
     "Runs one Network analysis at a time, because two Kilosort processes on one GPU "
     "exhaust its memory. Activity scans are CPU-only and run two at a time alongside. "
     "Folders can also be queued by hand from the control UI."),
    ("MEA-Analysis",
     "The lab's existing pipeline, launched unchanged: run_pipeline_driver starts one "
     "subprocess per well, which runs Kilosort4 on the GPU. The orchestrator chooses "
     "the interpreter and reads the driver's log to tell a real success from a run "
     "that exited zero having failed every well."),
    ("Activity scan",
     "Whole-array coverage before sorting: which electrodes are active, how much of "
     "the array is covered, and where the tissue sits. Kept separate from the network "
     "analysis because the two count different things."),
    ("Report builder",
     "Collects the analysed output and writes one report per session folder, into that "
     "folder. HTML for sending, PowerPoint for a lab meeting. A metric whose values "
     "cannot be right for its name is flagged at the top of the report rather than "
     "presented as fact."),
    ("Claude API (optional)",
     "Off unless enabled. Only computed summary statistics are sent — never recordings "
     "or raw data — and every figure in the returned prose is checked back against the "
     "measured values before the report may use it."),
    ("Control UI",
     "Reached from your own computer over an SSH tunnel; the server binds to localhost "
     "only. Folders, analyses, detection timing and the queue are all set here."),
]
for head, body in NOTES:
    p = d.add_paragraph()
    r = p.add_run(head + " — "); r.bold = True; r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string("2F6F52"); r.font.name = "Calibri"
    r2 = p.add_run(body); r2.font.size = Pt(10.5); r2.font.name = "Calibri"
    r2.font.color.rgb = RGBColor.from_string("3D4450")
    p.paragraph_format.space_after = Pt(8)

para("The recording drive is only ever read from, and the analysis code is never modified.",
     10, False, "6B7684", after=0)

d.save("/tmp/diag/RBSLab_Orchestrator_DataFlow.docx")
print("written")
